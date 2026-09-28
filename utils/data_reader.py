"""Reads test data from JSON or Excel (selected by DATA_SOURCE)."""
import json
import logging
import time

from openpyxl import load_workbook

from config import config

log = logging.getLogger(__name__)

REQUIRED_KEYS = ["firstname", "lastname", "email", "telephone",
                 "password", "product_name", "update_quantity"]


def _read_json(path):
    with open(path, encoding="utf-8") as fh:
        return json.load(fh)


def _read_excel(path):
    """Sheet 'TestData': row 1 = headers, row 2 = values."""
    wb = load_workbook(path, data_only=True)
    ws = wb["TestData"] if "TestData" in wb.sheetnames else wb.active
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 2:
        raise ValueError("Excel test data needs a header row and one data row.")
    headers = [str(h).strip() for h in rows[0] if h is not None]
    values = rows[1][: len(headers)]
    return dict(zip(headers, values))


def load_test_data(source: str = None) -> dict:
    source = (source or config.DATA_SOURCE).lower()
    if source == "json":
        data = _read_json(config.JSON_DATA_FILE)
    elif source in ("excel", "xlsx"):
        data = _read_excel(config.EXCEL_DATA_FILE)
    else:
        raise ValueError("DATA_SOURCE must be 'json' or 'excel'")

    missing = [k for k in REQUIRED_KEYS if data.get(k) in (None, "")]
    if missing:
        raise ValueError(f"Test data is missing values for: {missing}")

    data = {k: (str(v).strip() if not isinstance(v, (int, float)) else v)
            for k, v in data.items()}
    data["update_quantity"] = int(float(data["update_quantity"]))
    data["telephone"] = str(data["telephone"]).split(".")[0]
    # '{ts}' in the e-mail is replaced with a timestamp -> a unique user every run
    data["email"] = data["email"].replace("{ts}", str(int(time.time())))
    log.info("Loaded test data from %s (user=%s, product=%s, qty=%s)",
             source, data["email"], data["product_name"], data["update_quantity"])
    return data
